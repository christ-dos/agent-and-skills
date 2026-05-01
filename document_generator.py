"""Document generation module following the voyage-planning skill format."""

import re
import unicodedata
import datetime as _dt
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Any


_JOURS_FR = {
    0: 'LUNDI', 1: 'MARDI', 2: 'MERCREDI', 3: 'JEUDI',
    4: 'VENDREDI', 5: 'SAMEDI', 6: 'DIMANCHE',
}
_DAY_NAMES_PATTERN = '|'.join(_JOURS_FR.values())

# Format 1 : "VENDREDI 14/05" ou "VENDREDI 14/05/2026"
_DAY_BEFORE_RE = re.compile(
    rf'\b({_DAY_NAMES_PATTERN})\s+(\d{{1,2}})/(\d{{1,2}})(?:/\d{{2,4}})?\b',
    re.IGNORECASE,
)
# Format 2 : "14/05/2026 (Vendredi)" ou "14/05 (Vendredi)" — Mistral
_DAY_AFTER_RE = re.compile(
    rf'\b(\d{{1,2}})/(\d{{1,2}})(?:/\d{{2,4}})?\s*\(({_DAY_NAMES_PATTERN})\)',
    re.IGNORECASE,
)


def fix_day_names(text: str, year: int = None) -> str:
    """Replace incorrect day-of-week names with the real ones for the given year."""
    if year is None:
        m = re.search(r'\b(202[3-9]|203\d)\b', text)
        year = int(m.group(1)) if m else _dt.date.today().year

    def _correct(day_num, month_num):
        try:
            return _JOURS_FR[_dt.date(year, month_num, day_num).weekday()]
        except ValueError:
            return None

    def _replace_before(m):
        # "VENDREDI 14/05" → groupe 1=jour, 2=day, 3=month
        correct = _correct(int(m.group(2)), int(m.group(3)))
        if correct is None:
            return m.group(0)
        return m.group(0).replace(m.group(1), correct, 1)

    def _replace_after(m):
        # "14/05/2026 (Vendredi)" → groupe 1=day, 2=month, 3=jour
        correct = _correct(int(m.group(1)), int(m.group(2)))
        if correct is None:
            return m.group(0)
        # Preserve original capitalisation style of the day name
        wrong = m.group(3)
        replacement = correct.capitalize() if wrong[0].islower() or wrong[0].isupper() and wrong[1:].islower() else correct
        return m.group(0).replace(wrong, replacement, 1)

    text = _DAY_BEFORE_RE.sub(_replace_before, text)
    text = _DAY_AFTER_RE.sub(_replace_after, text)
    return text


DESTINATION_COLORS = {
    # Italie et villes
    'italie': {'primary': '8B1A1A', 'secondary': '2E6B3E', 'accent': 'C8961E', 'bg': 'FDF3E3'},
    'venise': {'primary': '8B1A1A', 'secondary': '2E6B3E', 'accent': 'C8961E', 'bg': 'FDF3E3'},
    'venice': {'primary': '8B1A1A', 'secondary': '2E6B3E', 'accent': 'C8961E', 'bg': 'FDF3E3'},
    'rome': {'primary': '8B1A1A', 'secondary': '2E6B3E', 'accent': 'C8961E', 'bg': 'FDF3E3'},
    'florence': {'primary': '8B1A1A', 'secondary': '2E6B3E', 'accent': 'C8961E', 'bg': 'FDF3E3'},
    'milan': {'primary': '8B1A1A', 'secondary': '2E6B3E', 'accent': 'C8961E', 'bg': 'FDF3E3'},
    # France et villes
    'france': {'primary': '003189', 'secondary': 'C41E3A', 'accent': 'D4AF37', 'bg': 'F0F4FF'},
    'paris': {'primary': '003189', 'secondary': 'C41E3A', 'accent': 'D4AF37', 'bg': 'F0F4FF'},
    'lyon': {'primary': '003189', 'secondary': 'C41E3A', 'accent': 'D4AF37', 'bg': 'F0F4FF'},
    'bordeaux': {'primary': '003189', 'secondary': 'C41E3A', 'accent': 'D4AF37', 'bg': 'F0F4FF'},
    # Japon et villes
    'japon': {'primary': 'C62A2F', 'secondary': '1A3A5C', 'accent': 'E8A800', 'bg': 'FFF5F5'},
    'tokyo': {'primary': 'C62A2F', 'secondary': '1A3A5C', 'accent': 'E8A800', 'bg': 'FFF5F5'},
    'kyoto': {'primary': 'C62A2F', 'secondary': '1A3A5C', 'accent': 'E8A800', 'bg': 'FFF5F5'},
    'osaka': {'primary': 'C62A2F', 'secondary': '1A3A5C', 'accent': 'E8A800', 'bg': 'FFF5F5'},
    # Grèce et villes
    'grece': {'primary': '1B4F8A', 'secondary': '1B8A6B', 'accent': 'E8C840', 'bg': 'F0F7FF'},
    'athenes': {'primary': '1B4F8A', 'secondary': '1B8A6B', 'accent': 'E8C840', 'bg': 'F0F7FF'},
    'santorini': {'primary': '1B4F8A', 'secondary': '1B8A6B', 'accent': 'E8C840', 'bg': 'F0F7FF'},
    'mykonos': {'primary': '1B4F8A', 'secondary': '1B8A6B', 'accent': 'E8C840', 'bg': 'F0F7FF'},
    # Espagne et villes
    'espagne': {'primary': 'C41E3A', 'secondary': 'F4A800', 'accent': '8B1A1A', 'bg': 'FFF8F0'},
    'madrid': {'primary': 'C41E3A', 'secondary': 'F4A800', 'accent': '8B1A1A', 'bg': 'FFF8F0'},
    'barcelone': {'primary': 'C41E3A', 'secondary': 'F4A800', 'accent': '8B1A1A', 'bg': 'FFF8F0'},
    'seville': {'primary': 'C41E3A', 'secondary': 'F4A800', 'accent': '8B1A1A', 'bg': 'FFF8F0'},
    # Portugal et villes
    'portugal': {'primary': '006B3C', 'secondary': 'C41E3A', 'accent': 'D4AF37', 'bg': 'F5FFF5'},
    'lisbonne': {'primary': '006B3C', 'secondary': 'C41E3A', 'accent': 'D4AF37', 'bg': 'F5FFF5'},
    'porto': {'primary': '006B3C', 'secondary': 'C41E3A', 'accent': 'D4AF37', 'bg': 'F5FFF5'},
    # Maroc et villes
    'maroc': {'primary': 'C8601A', 'secondary': '2A6B3A', 'accent': 'D4AF37', 'bg': 'FFF5EB'},
    'marrakech': {'primary': 'C8601A', 'secondary': '2A6B3A', 'accent': 'D4AF37', 'bg': 'FFF5EB'},
    'fes': {'primary': 'C8601A', 'secondary': '2A6B3A', 'accent': 'D4AF37', 'bg': 'FFF5EB'},
    # Royaume-Uni et villes
    'london': {'primary': '003580', 'secondary': '8B1A1A', 'accent': 'D4AF37', 'bg': 'F0F4FF'},
    'londres': {'primary': '003580', 'secondary': '8B1A1A', 'accent': 'D4AF37', 'bg': 'F0F4FF'},
    'edinburgh': {'primary': '003580', 'secondary': '8B1A1A', 'accent': 'D4AF37', 'bg': 'F0F4FF'},
    # Allemagne et villes
    'berlin': {'primary': '1A1A1A', 'secondary': 'C41E3A', 'accent': 'F4C800', 'bg': 'F8F8F8'},
    'munich': {'primary': '1A1A1A', 'secondary': 'C41E3A', 'accent': 'F4C800', 'bg': 'F8F8F8'},
}

DEFAULT_COLORS = {'primary': '1A376C', 'secondary': '007A87', 'accent': 'E8A800', 'bg': 'EBF5FB'}

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
URL_RE = re.compile(r'(https?://[^\s\)>\]\,\|«»]+)')

PPR_ORDER = [
    'pStyle', 'keepNext', 'keepLines', 'pageBreakBefore', 'framePr',
    'suppressLineNumbers', 'pBdr', 'shd', 'tabs', 'suppressAutoHyphens',
    'kinsoku', 'wordWrap', 'overflowPunct', 'topLinePunct', 'autoSpaceDE',
    'autoSpaceDN', 'bidi', 'adjustRightInd', 'snapToGrid', 'spacing', 'ind',
    'contextualSpacing', 'mirrorIndents', 'suppressOverlap', 'jc',
    'textDirection', 'textAlignment', 'textboxTightWrap', 'outlineLvl',
    'divId', 'cnfStyle', 'rPr', 'sectPr'
]


def extract_plan_from_conversation(conversation_history: List[Dict]) -> str:
    """Find the most recent assistant message that contains a planning block.

    Scans backwards, returns text trimmed to start at 'VOYAGE À' so that
    any short preamble ('Je vais maintenant…') is stripped.
    Falls back to the longest assistant text if no planning marker is found.
    """
    voyage_re = re.compile(r'VOYAGE\s+[AÀ]\s+', re.IGNORECASE)
    fallback = ""

    def _extract_text(msg) -> str:
        content = msg.get("content")
        if isinstance(content, list):
            for block in content:
                t = ""
                if hasattr(block, 'text') and getattr(block, 'type', None) == 'text':
                    t = block.text
                elif isinstance(block, dict) and block.get('type') == 'text':
                    t = block.get('text', '')
                if t.strip():
                    return t.strip()
            return ""
        if isinstance(content, str):
            return content.strip()
        return ""

    for msg in reversed(conversation_history):
        if msg.get("role") != "assistant":
            continue
        text = _extract_text(msg)
        if not text:
            continue

        m = voyage_re.search(text)
        if m:
            # Trim preamble — return from VOYAGE À onward
            return text[m.start():].strip()

        # Keep the longest non-planning text as fallback
        if len(text) > len(fallback):
            fallback = text

    return fallback


class DocumentGenerator:
    """Generate travel planning documents following the voyage-planning skill."""

    def __init__(self):
        Path("output").mkdir(exist_ok=True)

    def _slugify(self, text: str) -> str:
        """Convert destination name to slug: lowercase, no accents, hyphens."""
        text = text.lower().strip()
        text = unicodedata.normalize('NFD', text)
        text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
        text = re.sub(r'[^a-z0-9\-]', '-', text)
        text = re.sub(r'-+', '-', text).strip('-')
        return text

    def _get_colors(self, destination: str) -> dict:
        """Return color palette for the destination."""
        slug = self._slugify(destination)
        if slug in DESTINATION_COLORS:
            return DESTINATION_COLORS[slug]
        for key, colors in DESTINATION_COLORS.items():
            if key in slug or slug in key:
                return colors
        return DEFAULT_COLORS

    def _hex_to_rgb(self, hex_color: str):
        from docx.shared import RGBColor
        h = hex_color.lstrip('#')
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def _apply_xml_corrections(self, doc) -> None:
        """Apply mandatory XML corrections per the skill spec (inline, no unpack/pack)."""
        def ppr_sort_key(elem):
            tag = elem.tag.split('}')[-1] if '}' in elem.tag else elem.tag
            try:
                return PPR_ORDER.index(tag)
            except ValueError:
                return 999

        # Fix document.xml: sort pPr children + add missing shd val attribute
        for pPr in doc.element.iter(f'{{{W_NS}}}pPr'):
            children = list(pPr)
            sorted_children = sorted(children, key=ppr_sort_key)
            if [c.tag for c in children] != [c.tag for c in sorted_children]:
                for child in children:
                    pPr.remove(child)
                for child in sorted_children:
                    pPr.append(child)

        for shd in doc.element.iter(f'{{{W_NS}}}shd'):
            if shd.get(f'{{{W_NS}}}val') is None:
                shd.set(f'{{{W_NS}}}val', 'clear')

        # Fix settings.xml: replace w:val with w:percent on zoom element
        try:
            settings_elem = doc.part.settings.element
            for zoom in settings_elem.iter(f'{{{W_NS}}}zoom'):
                val_attr = f'{{{W_NS}}}val'
                pct_attr = f'{{{W_NS}}}percent'
                if zoom.get(val_attr) is not None:
                    val = zoom.attrib.pop(val_attr)
                    zoom.set(pct_attr, val)
        except Exception:
            pass

    def generate_planning_doc(self, planning_text: str, destination: str, step: int = 1) -> str:
        """
        Generate a Word document following the skill's 2-step structure.
        step=1 → planning initial (awaiting validation)
        step=2 → enriched planning with prices and booking links
        """
        dest_slug = self._slugify(destination)
        output_dir = Path("output") / dest_slug
        output_dir.mkdir(parents=True, exist_ok=True)

        base_name = f"planning_{dest_slug}" if step == 1 else f"planning_{dest_slug}_complet"
        # Avoid overwriting existing files: append numeric suffix if needed
        candidate = output_dir / f"{base_name}.docx"
        counter = 1
        while candidate.exists():
            candidate = output_dir / f"{base_name}_{counter}.docx"
            counter += 1
        doc_path = candidate

        md_path = output_dir / "planning.md"
        plan_text_final = planning_text

        # Toujours prioritiser le planning.md (source de vérité du planning validé)
        if md_path.exists():
            try:
                md_content = md_path.read_text(encoding="utf-8")
                m = re.search(r'## Plan détaillé\s*\n(.*?)(?=\n## |\Z)', md_content, re.S)
                if m:
                    section = m.group(1).strip()
                    if section:
                        plan_text_final = section
            except Exception:
                pass

        plan_text_final = fix_day_names(plan_text_final)
        colors = self._get_colors(destination)
        self._build_word_doc(plan_text_final, destination, colors, doc_path, step)
        return str(doc_path)

    def _build_word_doc(self, text: str, destination: str, colors: dict, path: Path, step: int) -> None:
        """Build and save the Word document with skill-compliant formatting."""
        try:
            from docx import Document
            from docx.shared import Pt, Inches, Cm
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            section = doc.sections[0]
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)

            normal_style = doc.styles['Normal']
            normal_style.font.name = 'Calibri'
            normal_style.font.size = Pt(11)

            primary = self._hex_to_rgb(colors['primary'])
            secondary = self._hex_to_rgb(colors['secondary'])

            # Main title
            title_para = doc.add_paragraph()
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_run = title_para.add_run(f"VOYAGE À {destination.upper()}")
            title_run.bold = True
            title_run.font.size = Pt(22)
            title_run.font.color.rgb = primary

            # Step 2 validation banner
            if step == 2:
                status_para = doc.add_paragraph()
                status_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                status_run = status_para.add_run("✅ Planning validé – Prix & réservations enrichis")
                status_run.italic = True
                status_run.font.size = Pt(11)
                status_run.font.color.rgb = secondary

            doc.add_paragraph()

            if text and text.strip():
                self._render_content(doc, text, colors, primary, secondary)
            else:
                p = doc.add_paragraph("Aucun planning généré.")
                p.runs[0].italic = True

            # Footer note per step
            doc.add_paragraph()
            if step == 1:
                note_para = doc.add_paragraph()
                note_run = note_para.add_run(
                    "Ce planning est en attente de validation. "
                    "Une fois validé, je lancerai les recherches de prix et de réservation."
                )
                note_run.italic = True
                note_run.font.size = Pt(10)

            doc.add_paragraph()
            footer_para = doc.add_paragraph()
            footer_run = footer_para.add_run(
                f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')} "
                "par l'Agent de Planification de Voyage"
            )
            footer_run.font.size = Pt(9)
            footer_run.italic = True

            # Page numbers in footer
            self._add_page_numbers(doc)

            # Apply mandatory XML corrections per the skill spec
            self._apply_xml_corrections(doc)

            doc.save(str(path))

        except ImportError:
            # Fallback: plain text if python-docx not installed
            txt_path = path.with_suffix('.txt')
            with open(txt_path, 'w', encoding='utf-8') as f:
                f.write(f"VOYAGE À {destination.upper()}\n")
                f.write(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}\n\n")
                f.write(text)

    def _render_content(self, doc, text: str, colors: dict, primary, secondary) -> None:
        """Parse and render AI response following the SKILL.md template."""
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        gray = RGBColor(0x66, 0x66, 0x66)

        day_re = re.compile(
            r'^(LUNDI|MARDI|MERCREDI|JEUDI|VENDREDI|SAMEDI|DIMANCHE|JOUR\s*\d+)',
            re.IGNORECASE
        )
        sep_re  = re.compile(r'^[═─=\-]{4,}$')
        title_re = re.compile(r'^VOYAGE\s+[AÀ]\s+', re.IGNORECASE)

        lines = text.split('\n')
        i = 0
        in_info_banner = False
        info_lines: list = []
        in_pensebete = False
        pensebete_lines: list = []

        def flush_info_banner():
            nonlocal in_info_banner, info_lines
            if info_lines:
                para = doc.add_paragraph()
                self._write_text_with_links(para, '  ·  '.join(info_lines))
                # Light background shading
                pPr = para._p.get_or_add_pPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), colors.get('bg', 'EBF5FB'))
                pPr.append(shd)
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.space_after  = Pt(8)
            in_info_banner = False
            info_lines = []

        def flush_pensebete():
            nonlocal in_pensebete, pensebete_lines
            if pensebete_lines:
                self._render_pensebete_box(doc, pensebete_lines)
            in_pensebete = False
            pensebete_lines = []

        while i < len(lines):
            raw     = lines[i]
            stripped = raw.strip()
            i += 1

            # ── empty line ──────────────────────────────────────────
            if not stripped:
                if in_info_banner or in_pensebete:
                    pass  # absorb blanks inside banner / pense-bête
                else:
                    doc.add_paragraph()
                continue

            # ── pense-bête header (must precede sep_re check) ───────
            if 'PENSE-BÊTE' in stripped.upper() or 'PENSE-BETE' in stripped.upper():
                in_pensebete = True
                pensebete_lines = []
                continue

            # ── pense-bête content / close ──────────────────────────
            if in_pensebete:
                if sep_re.match(stripped):
                    flush_pensebete()
                else:
                    pensebete_lines.append(stripped)
                continue

            # ── section separators (═══ / ───) ─────────────────────
            if sep_re.match(stripped):
                if in_info_banner:
                    flush_info_banner()
                # skip decorative separators
                continue

            # ── info banner header ──────────────────────────────────
            if 'Infos voyage' in stripped or stripped.startswith('─── Infos'):
                in_info_banner = True
                info_lines = []
                continue

            if in_info_banner:
                # any non-separator content is part of the banner
                info_lines.append(stripped)
                continue

            # ── main title (skip — already added in header) ─────────
            if title_re.match(stripped):
                continue

            # ── subtitle line: date | people | days ────────────────
            if re.match(r'.+–.+\|.+', stripped) and not stripped.startswith('|'):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(stripped)
                run.italic = True
                run.font.size = Pt(11)
                run.font.color.rgb = secondary
                continue

            # ── markdown headings ───────────────────────────────────
            if stripped.startswith('### '):
                p = doc.add_heading(stripped[4:].strip(), level=3)
                continue
            if stripped.startswith('## '):
                p = doc.add_heading(stripped[3:].strip(), level=2)
                if p.runs:
                    p.runs[0].font.color.rgb = secondary
                continue
            if stripped.startswith('# '):
                p = doc.add_heading(stripped[2:].strip(), level=1)
                if p.runs:
                    p.runs[0].font.color.rgb = primary
                continue

            # ── day headers (SAMEDI 04/04 – Thème / JOUR 1 – ...) ──
            if day_re.match(stripped):
                p = doc.add_heading(stripped, level=2)
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after  = Pt(4)
                if p.runs:
                    p.runs[0].font.color.rgb = secondary
                    p.runs[0].font.bold = True
                continue

            # ── section title lines (TABLEAU RÉCAPITULATIF / BUDGET) ─
            if re.match(r'^[A-ZÀÂÉÈÊËÎÏÔÙÛÜ\s&]{6,}$', stripped):
                p = doc.add_heading(stripped, level=2)
                if p.runs:
                    p.runs[0].font.color.rgb = primary
                continue

            # ── markdown table: collect consecutive | lines ─────────
            if stripped.startswith('|'):
                table_lines = [stripped]
                while i < len(lines) and lines[i].strip().startswith('|'):
                    table_lines.append(lines[i].strip())
                    i += 1
                self._render_md_table(doc, table_lines, colors)
                continue

            # ── sub-note lines (→ indented under an activity) ───────
            if stripped.startswith(('→', '->')):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.55)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after  = Pt(2)
                self._write_text_with_links(p, stripped)
                for run in p.runs:
                    run.font.color.rgb = gray
                    run.font.size = Pt(10)
                continue

            # ── activity bullets (•, ·) ─────────────────────────────
            if stripped.startswith(('•', '·')):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after  = Pt(2)
                self._write_text_with_links(p, stripped)
                continue

            # ── plain dash list items ───────────────────────────────
            if re.match(r'^[-–]\s', stripped):
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.left_indent = Inches(0.25)
                self._write_text_with_links(p, stripped[2:].strip())
                continue

            # ── default paragraph ───────────────────────────────────
            p = doc.add_paragraph()
            self._write_text_with_links(p, stripped)

        # flush any unclosed banner / pense-bête
        if in_info_banner:
            flush_info_banner()
        if in_pensebete:
            flush_pensebete()

    def _render_pensebete_box(self, doc, lines: list) -> None:
        """Render pense-bête lines as a yellow-background single-cell table."""
        from docx.shared import Pt, Inches
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        table = doc.add_table(rows=1, cols=1)
        table.style = 'Table Grid'
        cell = table.rows[0].cells[0]

        # Yellow fill on the cell
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'FFF8E1')
        tcPr.append(shd)

        # Clear the default empty paragraph
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)

        for line in lines:
            if not line:
                continue
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            self._write_text_with_links(p, line)
            for run in p.runs:
                run.font.size = Pt(10)

        doc.add_paragraph()  # spacing after box

    def _render_md_table(self, doc, table_lines: list, colors: dict) -> None:
        """Convert markdown table lines into a styled Word table."""
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        sep_re = re.compile(r'^\|[\s\-:|]+\|$')

        rows = []
        for line in table_lines:
            if sep_re.match(line):
                continue
            cells = [c.strip() for c in line.strip('|').split('|')]
            rows.append(cells)

        if not rows:
            return

        ncols = max(len(r) for r in rows)
        # Pad short rows
        rows = [r + [''] * (ncols - len(r)) for r in rows]

        table = doc.add_table(rows=len(rows), cols=ncols)
        table.style = 'Table Grid'

        for r_idx, row_data in enumerate(rows):
            tr = table.rows[r_idx]
            is_header = r_idx == 0
            is_total  = any('TOTAL' in c.upper() for c in row_data)

            for c_idx, cell_text in enumerate(row_data):
                cell = tr.cells[c_idx]
                cell.text = ''
                para = cell.paragraphs[0]
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
                self._write_text_with_links(para, cell_text)

                for run in para.runs:
                    if is_header:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(10)
                    elif is_total:
                        run.bold = True
                        run.font.size = Pt(10)

                # Header row background
                if is_header:
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'), colors['primary'])
                    tcPr.append(shd)
                elif is_total:
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'), colors.get('accent', 'E8A800'))
                    tcPr.append(shd)

        doc.add_paragraph()  # spacing after table

    def _add_hyperlink(self, para, url: str, display_text: str) -> None:
        """Insert a clickable hyperlink into a paragraph."""
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        HYPERLINK_NS = (
            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'
        )
        try:
            r_id = para.part.relate_to(url, HYPERLINK_NS, is_external=True)
        except Exception:
            para.add_run(display_text)
            return

        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), r_id)

        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')

        color = OxmlElement('w:color')
        color.set(qn('w:val'), '0563C1')
        rPr.append(color)

        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)

        r.append(rPr)

        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = display_text
        r.append(t)

        hyperlink.append(r)
        para._p.append(hyperlink)

    def _write_text_with_links(self, para, text: str) -> None:
        """Write text into a paragraph, turning http(s) URLs into clickable hyperlinks."""
        parts = URL_RE.split(text)
        for part in parts:
            if URL_RE.fullmatch(part):
                self._add_hyperlink(para, part, part)
            elif part:
                para.add_run(part)

    def _add_page_numbers(self, doc) -> None:
        """Add centered 'Page X / Y' numbers in the document footer."""
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        section = doc.sections[0]
        footer = section.footer
        para = footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        def append_field(para_elem, instr: str):
            for ftype, text in [
                ('begin', None),
                ('instr', instr),
                ('separate', None),
                ('end', None),
            ]:
                r = OxmlElement('w:r')
                rpr = OxmlElement('w:rPr')
                sz = OxmlElement('w:sz')
                sz.set(qn('w:val'), '18')
                rpr.append(sz)
                r.append(rpr)
                if ftype == 'instr':
                    it = OxmlElement('w:instrText')
                    it.set(qn('xml:space'), 'preserve')
                    it.text = f' {text} '
                    r.append(it)
                else:
                    fc = OxmlElement('w:fldChar')
                    fc.set(qn('w:fldCharType'), ftype if ftype != 'instr' else 'separate')
                    r.append(fc)
                para_elem.append(r)

        p = para._p

        def text_run(label):
            r = OxmlElement('w:r')
            rpr = OxmlElement('w:rPr')
            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), '18')
            rpr.append(sz)
            r.append(rpr)
            t = OxmlElement('w:t')
            t.set(qn('xml:space'), 'preserve')
            t.text = label
            r.append(t)
            p.append(r)

        text_run('Page ')
        append_field(p, 'PAGE')
        text_run(' / ')
        append_field(p, 'NUMPAGES')

    def generate_all_formats(
        self,
        conversation_history: List[Dict[str, Any]],
        destination: str = "voyage",
        step: int = 2,
    ) -> Dict[str, str]:
        """Generate doc from the last assistant message (complete planning)."""
        plan_text = extract_plan_from_conversation(conversation_history)
        doc_path = self.generate_planning_doc(plan_text, destination, step=step)
        return {"docx": doc_path}
